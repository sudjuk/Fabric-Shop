from PySide6.QtWidgets import QMessageBox, QPushButton, QFileDialog, QTableWidget
import psycopg2
import openpyxl

TABLES_CONFIG = {
    'employees': {
        'title': 'Сотрудники',
        'headers': ["ФИО", "Должность", "Телефон", "E-mail", "Дата приёма"],
        'fields': [
            {'name': 'full_name', 'label': 'ФИО', 'type': 'str'},
            {'name': 'position', 'label': 'Должность', 'type': 'str'},
            {'name': 'phone', 'label': 'Телефон', 'type': 'str'},
            {'name': 'email', 'label': 'E-mail', 'type': 'str'},
            {'name': 'hire_date', 'label': 'Дата приёма', 'type': 'date'},
        ],
        'pk': 'employee_id',
        'table': 'employees',
    },
    'roles': {
        'title': 'Роли',
        'headers': ["Название роли", "Пароль", "Описание"],
        'fields': [
            {'name': 'role_name', 'label': 'Название роли', 'type': 'str'},
            {'name': 'password', 'label': 'Пароль', 'type': 'password'},
            {'name': 'description', 'label': 'Описание', 'type': 'str'},
        ],
        'pk': 'role_id',
        'table': 'roles',
    },
    'fabrics': {
        'title': 'Ткани',
        'headers': ["Название", "Категория", "Цвет", "Плотность", "Ширина (см)", "Цена за м", "Остаток (м)", "Поставщик", "Дата поступления"],
        'fields': [
            {'name': 'name', 'label': 'Название', 'type': 'str'},
            {'name': 'category_id', 'label': 'Категория', 'type': 'combo', 'combo_table': 'fabric_categories', 'combo_label': 'name'},
            {'name': 'color', 'label': 'Цвет', 'type': 'str'},
            {'name': 'density', 'label': 'Плотность', 'type': 'str'},
            {'name': 'width_cm', 'label': 'Ширина (см)', 'type': 'float'},
            {'name': 'price_per_meter', 'label': 'Цена за м', 'type': 'float'},
            {'name': 'stock_meters', 'label': 'Остаток (м)', 'type': 'float'},
            {'name': 'supplier_id', 'label': 'Поставщик', 'type': 'combo', 'combo_table': 'suppliers', 'combo_label': 'company_name'},
            {'name': 'arrival_date', 'label': 'Дата поступления', 'type': 'date'},
        ],
        'pk': 'fabric_id',
        'table': 'fabrics',
    },
    'fabric_categories': {
        'title': 'Категории тканей',
        'headers': ["Название", "Описание"],
        'fields': [
            {'name': 'name', 'label': 'Название', 'type': 'str'},
            {'name': 'description', 'label': 'Описание', 'type': 'str'},
        ],
        'pk': 'category_id',
        'table': 'fabric_categories',
    },
    'suppliers': {
        'title': 'Поставщики',
        'headers': ["Компания", "Контактное лицо", "Телефон", "E-mail", "Адрес"],
        'fields': [
            {'name': 'company_name', 'label': 'Компания', 'type': 'str'},
            {'name': 'contact_person', 'label': 'Контактное лицо', 'type': 'str'},
            {'name': 'phone', 'label': 'Телефон', 'type': 'str'},
            {'name': 'email', 'label': 'E-mail', 'type': 'str'},
            {'name': 'address', 'label': 'Адрес', 'type': 'str'},
        ],
        'pk': 'supplier_id',
        'table': 'suppliers',
    },
    'customers': {
        'title': 'Клиенты',
        'headers': ["ФИО", "Телефон", "E-mail", "Адрес"],
        'fields': [
            {'name': 'full_name', 'label': 'ФИО', 'type': 'str'},
            {'name': 'phone', 'label': 'Телефон', 'type': 'str'},
            {'name': 'email', 'label': 'E-mail', 'type': 'str'},
            {'name': 'address', 'label': 'Адрес', 'type': 'str'},
        ],
        'pk': 'customer_id',
        'table': 'customers',
    },
    'deliveries': {
        'title': 'Поставки',
        'headers': ["Дата", "Поставщик", "Сотрудник"],
        'fields': [
            {'name': 'delivery_date', 'label': 'Дата', 'type': 'date'},
            {'name': 'supplier_id', 'label': 'Поставщик', 'type': 'combo', 'combo_table': 'suppliers', 'combo_label': 'company_name'},
            {'name': 'employee_id', 'label': 'Сотрудник', 'type': 'combo', 'combo_table': 'employees', 'combo_label': 'full_name'},
        ],
        'pk': 'delivery_id',
        'table': 'deliveries',
    },
    'sales': {
        'title': 'Продажи',
        'headers': ["Дата продажи", "Сотрудник", "Клиент", "Сумма", "Способ оплаты", "Примечание"],
        'fields': [
            {'name': 'sale_date', 'label': 'Дата продажи', 'type': 'date'},
            {'name': 'employee_id', 'label': 'Сотрудник', 'type': 'combo', 'combo_table': 'employees', 'combo_label': 'full_name'},
            {'name': 'customer_id', 'label': 'Клиент', 'type': 'combo', 'combo_table': 'customers', 'combo_label': 'full_name'},
            {'name': 'total_amount', 'label': 'Сумма', 'type': 'float'},
            {'name': 'payment_method', 'label': 'Способ оплаты', 'type': 'str'},
            {'name': 'notes', 'label': 'Примечание', 'type': 'str'},
        ],
        'pk': 'sale_id',
        'table': 'sales',
    },
    'sale_items': {
        'title': 'Содержание продажи',
        'headers': ["Продажа", "Ткань", "Количество (м)", "Цена за м", "Сумма"],
        'fields': [
            {'name': 'sale_id', 'label': 'Продажа', 'type': 'combo', 'combo_table': 'sales', 'combo_label': 'sale_id'},
            {'name': 'fabric_id', 'label': 'Ткань', 'type': 'combo', 'combo_table': 'fabrics', 'combo_label': 'name'},
            {'name': 'quantity_meters', 'label': 'Количество (м)', 'type': 'float'},
            {'name': 'price_per_meter', 'label': 'Цена за м', 'type': 'float'},
            {'name': 'total_amount', 'label': 'Сумма', 'type': 'float'},
        ],
        'pk': 'sale_item_id',
        'table': 'sale_items',
    },
}

EXPORT_BTN_STYLE = """
    QPushButton {
        background: #0078d7;
        color: #fff;
        border-radius: 18px;
        font-size: 20px;
        font-weight: bold;
        border: none;
    }
    QPushButton:hover {
        background: #005fa3;
    }
"""

TABLE_STYLE = """
    QTableWidget { background: #232323; color: #fff; font-size: 15px; border-radius: 8px; }
    QHeaderView::section { background: #333; color: #fff; font-weight: bold; font-size: 16px; padding: 6px; border: none; }
    QTableWidget::item { padding: 6px; }
    QTableWidget QTableCornerButton::section { background: #333; border: none; }
"""

BUTTON_STYLE = """
    QPushButton { background: #0078d7; color: #fff; border-radius: 8px; padding: 8px 18px; font-size: 15px; font-weight: bold; }
    QPushButton:hover { background: #005fa3; }
"""

def get_db_connection():
    return psycopg2.connect(
        dbname="fabric_shop",
        user="postgres",
        password="lol132kek",
        host="localhost",
        port="5432"
    )

def get_combo_values(table, label_field):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        pk = TABLES_CONFIG[table]['pk'] if table in TABLES_CONFIG else table[:-1] + '_id'
        cursor.execute(f"SELECT {pk}, {label_field} FROM {table}")
        values = cursor.fetchall()
        cursor.close()
        conn.close()
        return values
    except psycopg2.Error as e:
        QMessageBox.critical(None, "Ошибка", f"Ошибка загрузки справочника: {e}")
        return []

def load_table_data(table_key):
    config = TABLES_CONFIG[table_key]
    pk = config['pk']
    display_fields = [f for f in config['fields']]
    select_fields = []
    joins = []
    for f in display_fields:
        if f.get('type') == 'combo':
            combo_table = f['combo_table']
            combo_label = f['combo_label']
            combo_pk = TABLES_CONFIG[combo_table]['pk'] if combo_table in TABLES_CONFIG else combo_table[:-1] + '_id'
            select_fields.append(f"{combo_table}.{combo_label} AS {f['name']}_label")
            joins.append(f"LEFT JOIN {combo_table} ON {config['table']}.{f['name']} = {combo_table}.{combo_pk}")
        else:
            select_fields.append(f"{config['table']}.{f['name']}")
    select_fields = [f"{config['table']}.{pk}"] + select_fields
    query = f"SELECT {', '.join(select_fields)} FROM {config['table']} {' '.join(joins)}"
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(query)
        rows = cursor.fetchall()
        cursor.close()
        conn.close()
        return rows
    except psycopg2.Error as e:
        QMessageBox.critical(None, "Ошибка", f"Ошибка загрузки данных: {e}")
        return []

def add_row_to_table(table_key, values):
    config = TABLES_CONFIG[table_key]
    field_names = [f['name'] for f in config['fields']]
    placeholders = ', '.join(['%s'] * len(field_names))
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            f"INSERT INTO {config['table']} ({', '.join(field_names)}) VALUES ({placeholders})",
            tuple(values[name] for name in field_names)
        )
        conn.commit()
        cursor.close()
        conn.close()
        return True
    except psycopg2.Error as e:
        QMessageBox.critical(None, "Ошибка", f"Ошибка добавления: {e}")
        return False

def update_row_in_table(table_key, values, pk_value):
    config = TABLES_CONFIG[table_key]
    pk = config['pk']
    set_expr = ', '.join([f"{f['name']}=%s" for f in config['fields']])
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            f"UPDATE {config['table']} SET {set_expr} WHERE {pk}=%s",
            tuple(values[name] for name in [f['name'] for f in config['fields']]) + (pk_value,)
        )
        conn.commit()
        cursor.close()
        conn.close()
        return True
    except psycopg2.Error as e:
        QMessageBox.critical(None, "Ошибка", f"Ошибка изменения: {e}")
        return False

def delete_row_from_table(table_key, pk_value):
    config = TABLES_CONFIG[table_key]
    pk = config['pk']
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(f"DELETE FROM {config['table']} WHERE {pk}=%s", (pk_value,))
        conn.commit()
        cursor.close()
        conn.close()
        return True
    except psycopg2.Error as e:
        QMessageBox.critical(None, "Ошибка", f"Ошибка удаления: {e}")
        return False

def create_export_button(parent, callback):
    btn = QPushButton("⇩", parent)
    btn.setToolTip("Экспорт в Excel")
    btn.setFixedSize(36, 36)
    btn.setStyleSheet(EXPORT_BTN_STYLE)
    btn.clicked.connect(callback)
    btn.raise_()
    return btn

def export_table_to_excel(table: QTableWidget, title: str):
    path, _ = QFileDialog.getSaveFileName(table, f"Сохранить отчёт в Excel", "", "Excel Files (*.xlsx)")
    if not path:
        return
    if not path.endswith('.xlsx'):
        path += '.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    headers = [table.horizontalHeaderItem(i).text() for i in range(table.columnCount())]
    ws.append(headers)
    for row in range(table.rowCount()):
        ws.append([
            table.item(row, col).text() if table.item(row, col) else ''
            for col in range(table.columnCount())
        ])
    wb.save(path)

def get_table_style():
    return TABLE_STYLE

def get_button_style():
    return BUTTON_STYLE 